#!/usr/bin/env python3
"""Build a stratified, non-copying manifest for angle evaluation."""

from __future__ import annotations

import argparse
import csv
import random
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

from PIL import Image, UnidentifiedImageError


TYPE_TO_VIEW = {
    1: "left_front",
    2: "right_front",
    3: "left_rear",
    4: "right_rear",
    10: "interior_front",
    11: "interior_rear",
}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}
NAME_RE = re.compile(r"^(\d{1,2})_[^_]+(?:_[^_]+)?__(.+)$")
MANIFEST_FIELDS = (
    "image_path",
    "order_number",
    "car_no",
    "image_type",
    "ground_truth_view",
    "original_filename",
)


def view_for_image_type(image_type: int | str) -> str | None:
    try:
        return TYPE_TO_VIEW.get(int(image_type))
    except (TypeError, ValueError):
        return None


def image_is_valid(path: Path) -> bool:
    try:
        with Image.open(path) as image:
            image.verify()
        return True
    except (OSError, UnidentifiedImageError):
        return False


def rows_from_directory(image_root: Path) -> Iterable[dict[str, str]]:
    for path in sorted(image_root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in IMAGE_SUFFIXES:
            continue
        relative = path.relative_to(image_root)
        if len(relative.parts) < 3:
            continue
        match = NAME_RE.match(path.name)
        if not match:
            continue
        image_type = int(match.group(1))
        view = view_for_image_type(image_type)
        if view is None:
            continue
        yield {
            "image_path": path.resolve().as_posix(),
            "order_number": relative.parts[-2],
            "car_no": relative.parts[-3],
            "image_type": str(image_type),
            "ground_truth_view": view,
            "original_filename": match.group(2),
        }


def rows_from_excel(excel_path: Path, image_root: Path) -> Iterable[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("使用 --source excel 需要 openpyxl") from error
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    required = {"order_number", "CarNo", "ImageType", "Image"}
    missing = required.difference(headers)
    if missing:
        raise ValueError(f"Excel 缺少欄位：{', '.join(sorted(missing))}")
    indexes = {name: headers.index(name) for name in required}
    for values in rows:
        image_type_raw = values[indexes["ImageType"]]
        view = view_for_image_type(image_type_raw)
        if view is None:
            continue
        image_type = int(image_type_raw)
        order = str(values[indexes["order_number"]]).strip()
        car_no = str(values[indexes["CarNo"]]).strip()
        original = str(values[indexes["Image"]]).strip()
        expected = image_root / car_no / order / f"{image_type:02d}_{view}__{original}"
        yield {
            "image_path": expected.resolve().as_posix(),
            "order_number": order,
            "car_no": car_no,
            "image_type": str(image_type),
            "ground_truth_view": view,
            "original_filename": original,
        }


def sample_manifest_rows(
    rows: Iterable[dict[str, str]], per_class: int = 20, seed: int = 42
) -> tuple[list[dict[str, str]], list[str]]:
    by_view_order: dict[str, dict[str, list[dict[str, str]]]] = defaultdict(lambda: defaultdict(list))
    invalid = 0
    for row in rows:
        path = Path(row["image_path"])
        if not path.is_file() or not image_is_valid(path):
            invalid += 1
            continue
        by_view_order[row["ground_truth_view"]][row["order_number"]].append(row)

    rng = random.Random(seed)
    sampled: list[dict[str, str]] = []
    warnings: list[str] = []
    if invalid:
        warnings.append(f"略過 {invalid} 張不存在或 Pillow 無法開啟的圖片")
    for view in TYPE_TO_VIEW.values():
        groups = by_view_order.get(view, {})
        orders = sorted(groups)
        rng.shuffle(orders)
        selected_orders = orders[:per_class]
        if len(selected_orders) < per_class:
            warnings.append(f"{view} 僅有 {len(selected_orders)} 筆有效且不同訂單的圖片，少於 {per_class}")
        for order in selected_orders:
            choices = sorted(groups[order], key=lambda row: row["image_path"])
            sampled.append(rng.choice(choices))
    return sampled, warnings


def write_manifest(rows: list[dict[str, str]], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=MANIFEST_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="建立六個主要角度的分層抽樣 manifest，不複製圖片")
    parser.add_argument("--image-root", type=Path, default=Path("downloaded_images"))
    parser.add_argument("--excel", type=Path, default=Path("車外車內照片上傳清單.xlsx"))
    parser.add_argument("--source", choices=("directory", "excel"), default="directory")
    parser.add_argument("--output", type=Path, default=Path("data/eval/angle_eval_manifest.csv"))
    parser.add_argument("--per-class", type=int, default=20)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    if args.per_class < 1:
        parser.error("--per-class 必須大於 0")
    if not args.image_root.is_dir():
        parser.error(f"找不到圖片資料夾：{args.image_root}")
    source_rows = (
        rows_from_excel(args.excel, args.image_root)
        if args.source == "excel"
        else rows_from_directory(args.image_root)
    )
    try:
        sampled, warnings = sample_manifest_rows(source_rows, args.per_class, args.seed)
    except (OSError, RuntimeError, ValueError) as error:
        print(f"錯誤：{error}", file=sys.stderr)
        return 1
    write_manifest(sampled, args.output)
    counts = Counter(row["ground_truth_view"] for row in sampled)
    for warning in warnings:
        print(f"警告：{warning}", file=sys.stderr)
    print(f"Manifest：{args.output}（共 {len(sampled)} 張）")
    for view in TYPE_TO_VIEW.values():
        print(f"{view}: {counts[view]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
